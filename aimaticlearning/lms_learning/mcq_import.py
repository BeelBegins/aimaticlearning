from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

import frappe
from docx import Document
from frappe import _
from frappe.utils import cint, get_site_path

from aimaticlearning.lms_learning.content_generation import _link_question_to_quiz
from aimaticlearning.lms_learning.import_pipeline import _ensure_chapter_quiz
from aimaticlearning.lms_learning.outline_sync import sync_profile_outline

DEFAULT_PER_CHAPTER_TARGET = 20
DEFAULT_MODULE_ASSESSMENT_TARGET = 150

QUESTION_RE = re.compile(r"^Question\s+(\d+(?:\.\d+)?)\s*[.)]?\s*(.*)$", re.I)
NUMBERED_QUESTION_RE = re.compile(r"^(\d+(?:\.\d+)?)[.)]\s+(.+)$")
ANSWER_RE = re.compile(
	r"^(?:Question\s+(\d+(?:\.\d+)?)\s*[:\-–]\s*(?:(?:Correct\s+)?Answer\s*[:\-–]?\s*)?"
	r"|Answer\s+(\d+(?:\.\d+)?)\s*[:.\)\-]\s*"
	r"|(\d+(?:\.\d+)?)\s*[.)]\s*)([A-E])(?=\s|[-–:]|$)"
	r"(?:\s*[-–:]\s*|\s+)?(.*)$",
	re.I,
)
BARE_ANSWER_RE = re.compile(r"^([A-E])(?:\s*[-–:]\s*|\s+)(.+)$", re.I)
CORRECT_ANSWER_RE = re.compile(r"^Correct Answer:\s*([A-E])\s*-\s*(.*)$", re.I)
EXPLANATION_RE = re.compile(r"^(?:Explanation|Reason)\s*:\s*(.*)$", re.I)
LEARNER_SOURCE_TRAIL_RE = re.compile(
	r"(?:\s|<br\s*/?>)*(?:source|source reference)\s*:\s*(?:file://)?(?:/home/|/private/files/|/files/).*$",
	re.I | re.S,
)
OPTION_SPLIT_RE = re.compile(r"\s+(?=[A-E][\.\)])")
OPTION_RE = re.compile(r"^([A-E])[\.\)]\s*(.*)$", re.I)
SECTION_RE = re.compile(
	r"(SQE1 Practice Questions|Multiple Choice Questions|MCQ Practice|Practice MCQs|"
	r"Practice Questions|MCQ Assessment|MCQ Practice Test|MCQs for Chapter|^Questions\s+\d)",
	re.I,
)
ANSWER_SECTION_RE = re.compile(
	r"^(?:answer key|answers?\s+(?:and|with|&)|correct answers?)", re.I
)

FLK2_MCQ_SPECS = {
	"Criminal Law Notes for SQE1.docx": {
		"subject": "Criminal Law",
		"question_counts": (20, 20, 20, 10, 20, 10),
	},
	"Criminal Litigation.docx": {
		"subject": "Criminal Litigation",
		"question_counts": (20,) * 11,
	},
	"Eq nd trust.docx": {
		"subject": "Equity and Trust Law",
		"question_counts": (20,) * 8,
	},
	"Land Law.docx": {
		"subject": "Land Law",
		"question_counts": (20,) * 8,
	},
	"Property Practice Essentials for SQE1.docx": {
		"subject": "Property Practice",
		# Latest red-key Word bank has 20 MCQs in every chapter (incl. Ch3 Planning, Ch7 Remedies).
		"question_counts": (20,) * 8,
	},
	"Solicitor ACC.docx": {
		"subject": "Solicitors' Accounts",
		"question_counts": (20,) * 8,
	},
	"WILLS AND THE ADMINISTRATION OF ESTATES.docx": {
		"subject": "Wills and Administration of Estates",
		"question_counts": (20,) * 6,
	},
}

FLK2_SECTION_RE = re.compile(
	r"(?:multiple\s+choice\s+questions|mcqs?\s+for\s+chapter|practice\s+mcqs?|"
	r"mcq\s+(?:assessment|practice\s+test|practice\s+questions?))",
	re.I,
)
FLK2_ANSWER_HEADING_RE = re.compile(
	r"^(?:answers?|answer\s+key|correct\s+answers?)(?:\b|\s*[:&-])", re.I
)
FLK2_QUESTION_RE = re.compile(r"^Question\s+(\d+(?:\.\d+)?)\s*(.*)$", re.I)
FLK2_NUMBERED_QUESTION_RE = re.compile(r"^(\d+(?:\.\d+)?)[.)]\s+(.+)$")
FLK2_EXPLICIT_ANSWER_RE = re.compile(
	r"^Question\s+(\d+(?:\.\d+)?)\s*[-:]\s*Answer\s*:?[ \t]*([A-E])"
	r"(?:\s*[-–:]\s*|\s+)?(.*)$",
	re.I,
)
FLK2_ANSWER_RE = re.compile(
	r"^Answer\s+(\d+(?:\.\d+)?)\s*[:.)-]\s*([A-E])"
	r"(?:\s*[-–:]\s*|\s+)?(.*)$",
	re.I,
)
FLK2_NUMBERED_ANSWER_RE = re.compile(
	r"^(\d+(?:\.\d+)?)\s*[-.)]\s*([A-E])(?:\s*[-–:]\s*|\s+)?(.*)$",
	re.I,
)
FLK2_BARE_ANSWER_RE = re.compile(r"^([A-E])\s*[-–:]\s*(.+)$", re.I)
FLK2_BARE_LETTER_RE = re.compile(r"^([A-E])$", re.I)
FLK2_FINAL_ANSWER_RE = re.compile(r"^Final\s+Answer\s*:", re.I)


def parse_mcqs_from_docx(path: Path, limit: int | None = None) -> list[dict]:
	document = Document(str(path))
	paragraphs = [" ".join((p.text or "").replace("\xa0", " ").split()) for p in document.paragraphs]
	source_label = path.name

	all_questions: list[dict] = []
	section_id = 0
	i = 0
	while i < len(paragraphs):
		text = paragraphs[i]
		if not text or not SECTION_RE.search(text):
			i += 1
			continue
		section_id += 1
		section_question_count = 0
		block_questions: dict[str, dict] = {}
		active_num: str | None = None
		i += 1
		# Question phase
		while i < len(paragraphs):
			text = paragraphs[i]
			if not text:
				i += 1
				continue
			if SECTION_RE.search(text):
				break
			if ANSWER_SECTION_RE.match(text) or ANSWER_RE.match(text):
				break
			qm = QUESTION_RE.match(text)
			if not qm:
				numbered = NUMBERED_QUESTION_RE.match(text)
				if numbered:
					qm = numbered
			if qm:
				num = qm.group(1)
				stem = qm.group(2).strip()
				if not stem and i + 1 < len(paragraphs):
					next_text = paragraphs[i + 1]
					if next_text and not QUESTION_RE.match(next_text) and not OPTION_RE.match(
						next_text
					):
						stem = next_text
						i += 1
				block_questions[num] = {
					"num": num,
					"question": stem or f"Question {num}",
					"options": [],
				}
				active_num = num
				i += 1
				continue
			if active_num:
				opts = _split_options(text)
				if opts:
					block_questions[active_num]["options"].extend(opts)
			i += 1

		# Answer phase (explicit key header or inline "Question N: Answer X")
		current_num: str | None = None
		answer_order = list(block_questions)
		answer_cursor = 0
		while i < len(paragraphs):
			text = paragraphs[i]
			if not text:
				i += 1
				continue
			if SECTION_RE.search(text):
				break
			if ANSWER_SECTION_RE.match(text):
				i += 1
				continue
			am = ANSWER_RE.match(text)
			if am:
				current_num = am.group(1) or am.group(2) or am.group(3)
				block_questions.setdefault(current_num, {"num": current_num})
				block_questions[current_num]["correct_letter"] = am.group(4).upper()
				summary = _without_explanation_label(am.group(5))
				if summary:
					_append_explanation(block_questions[current_num], summary)
				if current_num in answer_order:
					answer_cursor = answer_order.index(current_num) + 1
				i += 1
				continue
			cam = CORRECT_ANSWER_RE.match(text)
			if cam and current_num:
				block_questions[current_num]["correct_letter"] = cam.group(1).upper()
				_append_explanation(block_questions[current_num], cam.group(2))
				i += 1
				continue
			rm = EXPLANATION_RE.match(text)
			if rm and current_num:
				block_questions.setdefault(current_num, {"num": current_num})
				_append_explanation(block_questions[current_num], rm.group(1))
				i += 1
				continue
			bam = BARE_ANSWER_RE.match(text)
			if bam and answer_cursor < len(answer_order):
				current_num = answer_order[answer_cursor]
				block_questions[current_num]["correct_letter"] = bam.group(1).upper()
				_append_explanation(block_questions[current_num], bam.group(2))
				answer_cursor += 1
				i += 1
				continue
			if QUESTION_RE.match(text):
				break
			if current_num:
				_append_explanation(block_questions[current_num], text)
			i += 1

		for num in block_questions:
			row = block_questions[num]
			if not row.get("options"):
				continue
			normalized = _normalize_question_row(row, section_id, num, source_label)
			if not any(opt.get("is_correct") for opt in normalized["options"]):
				continue
			all_questions.append(normalized)
			section_question_count += 1
			if limit and len(all_questions) >= limit:
				return all_questions
		if not section_question_count:
			section_id -= 1

	return all_questions


def sync_mcq_explanations_from_source(source_file: str, dry_run: bool = True) -> dict:
	"""Apply vetted DOCX reasons only where the existing mapped MCQ still matches."""
	path = _resolve_source_path(source_file)
	parsed = parse_mcqs_from_docx(path)
	result = {
		"source_file": path.name,
		"parsed_questions": len(parsed),
		"matched_questions": 0,
		"ready_to_update": 0,
		"updated_questions": 0,
		"unmapped": [],
		"mismatched": [],
		"missing_source_reason": [],
	}
	for item in parsed:
		source_refs = [
			item.get("source_reference"),
			f"{path.name} Ch{item.get('section_id')} Q{item.get('num')}",
		]
		meta = None
		for source_ref in source_refs:
			if not source_ref:
				continue
			meta = frappe.db.get_value(
				"Learning Question Meta",
				{"source_reference": source_ref},
				["lms_question", "source_reference"],
				as_dict=True,
			)
			if meta:
				break
		if not meta or not meta.lms_question:
			result["unmapped"].append(item.get("source_reference"))
			continue

		doc = frappe.get_doc("LMS Question", meta.lms_question)
		expected_correct = [
			index
			for index, option in enumerate(item.get("options") or [], start=1)
			if option.get("is_correct")
		]
		stored_correct = [
			index for index in range(1, 6) if cint(doc.get(f"is_correct_{index}"))
		]
		if (
			normalize_text(doc.question) != normalize_text(item.get("question"))
			or stored_correct != expected_correct
		):
			result["mismatched"].append(meta.source_reference)
			continue

		result["matched_questions"] += 1
		correct_index = expected_correct[0] if len(expected_correct) == 1 else None
		reason = (
			item.get("options", [])[correct_index - 1].get("explanation")
			if correct_index
			else ""
		)
		reason = (reason or "").strip()
		if not reason:
			result["missing_source_reason"].append(meta.source_reference)
			continue

		field = f"explanation_{correct_index}"
		current = (doc.get(field) or "").strip()
		if current and current != reason:
			result["mismatched"].append(meta.source_reference)
			continue
		if current == reason:
			continue

		result["ready_to_update"] += 1
		if not dry_run:
			doc.set(field, reason)
			doc.save(ignore_permissions=True)
			result["updated_questions"] += 1

	if not dry_run:
		frappe.db.commit()
	return result


def normalize_text(value: str | None) -> str:
	return " ".join((value or "").split())


def parse_verified_flk2_mcqs_from_docx(path: Path | str) -> dict:
	"""Parse FLK2 answer keys with fail-closed section and question checks."""
	path = Path(path)
	spec = FLK2_MCQ_SPECS.get(path.name)
	if not spec:
		raise ValueError(f"No FLK2 MCQ specification for {path.name}")

	paragraphs = [" ".join((p.text or "").replace("\xa0", " ").split()) for p in Document(str(path)).paragraphs]
	starts = [
		index
		for index, text in enumerate(paragraphs)
		if _is_flk2_section_heading(text)
	]
	issues: list[str] = []
	expected_counts = tuple(spec["question_counts"])
	if len(starts) != len(expected_counts):
		issues.append(
			f"{path.name}: expected {len(expected_counts)} MCQ sections, found {len(starts)}"
		)

	sections = []
	all_questions = []
	for section_id, start in enumerate(starts, start=1):
		end = starts[section_id] if section_id < len(starts) else len(paragraphs)
		section_title = paragraphs[start]
		questions = {}
		question_order: list[str] = []
		answers = {}
		answer_heading = next(
			(index for index in range(start + 1, end) if FLK2_ANSWER_HEADING_RE.match(paragraphs[index])),
			None,
		)
		section_issues: list[str] = []
		if answer_heading is None:
			section_issues.append("missing answer heading")
		else:
			questions, question_order, question_issues = _parse_flk2_questions(
				paragraphs[start + 1 : answer_heading]
			)
			section_issues.extend(question_issues)
			answers, answer_issues = _parse_flk2_answers(
				paragraphs[answer_heading + 1 : end], question_order
			)
			section_issues.extend(answer_issues)
			section_issues.extend(_validate_flk2_section(questions, answers, question_order))

		expected = expected_counts[section_id - 1] if section_id <= len(expected_counts) else None
		if expected is not None and len(questions) != expected:
			section_issues.append(f"expected {expected} questions, found {len(questions)}")
		if expected is not None and len(answers) != expected:
			section_issues.append(f"expected {expected} answers, found {len(answers)}")

		section_row = {
			"section_id": section_id,
			"section_title": section_title,
			"paragraph_start": start,
			"paragraph_end": end - 1,
			"question_count": len(questions),
			"answer_count": len(answers),
			"issues": [f"S{section_id}: {issue}" for issue in section_issues],
		}
		sections.append(section_row)
		issues.extend(section_row["issues"])

		for num in question_order:
			if num not in questions or num not in answers:
				continue
			row = questions[num]
			row.update(answers[num])
			row["num"] = num
			row["section_id"] = section_id
			row["section_title"] = section_title
			row["source_reference"] = f"{path.name} S{section_id} Q{num}"
			row["options"] = [
				{
					"text": option["text"],
					"is_correct": option["letter"] == row.get("correct_letter"),
					"explanation": row.get("explanation", "")
					if option["letter"] == row.get("correct_letter")
					else "",
				}
				for option in questions[num]["options"]
			]
			all_questions.append(row)

	return {
		"ready": not issues,
		"subject": spec["subject"],
		"source_file": path.name,
		"sections": sections,
		"questions": all_questions,
		"question_count": len(all_questions),
		"issues": issues,
	}


def _is_flk2_section_heading(text: str) -> bool:
	if not text or not FLK2_SECTION_RE.search(text):
		return False
	lower = text.lower()
	return not any(marker in lower for marker in ("test score", "scoring guide", "self-assessment"))


def _split_flk2_inline_options(text: str, allow_non_a_start: bool = False) -> tuple[str, list[dict]]:
	markers = list(re.finditer(r"(?<!\S)([A-E])[.)]\s+", text, re.I))
	if not markers:
		return text.strip(), []
	marker_start = 0
	letters = [marker.group(1).upper() for marker in markers]
	if markers[0].start() != 0:
		marker_start = next(
			(
				index
				for index in range(len(markers) - 4)
				if letters[index : index + 5] == ["A", "B", "C", "D", "E"]
			),
			None,
		)
		if marker_start is None:
			return text.strip(), []
	elif letters[0] != "A" and not allow_non_a_start:
		return text.strip(), []
	markers = markers[marker_start:]
	if markers[0].start() == 0:
		selected = [markers[0]]
		expected_index = "ABCDE".find(markers[0].group(1).upper()) + 1
		for marker in markers[1:]:
			if expected_index >= len("ABCDE"):
				break
			if marker.group(1).upper() == "ABCDE"[expected_index]:
				selected.append(marker)
				expected_index += 1
		markers = selected
	options = []
	for index, marker in enumerate(markers):
		end = markers[index + 1].start() if index + 1 < len(markers) else len(text)
		options.append(
			{
				"letter": marker.group(1).upper(),
				"text": text[marker.end() : end].strip(),
			}
		)
	return text[: markers[0].start()].strip(), options


def _parse_flk2_questions(lines: list[str]) -> tuple[dict, list[str], list[str]]:
	questions: dict[str, dict] = {}
	question_order: list[str] = []
	issues: list[str] = []
	active: dict | None = None
	for text in lines:
		if not text:
			continue
		match = FLK2_QUESTION_RE.match(text) or FLK2_NUMBERED_QUESTION_RE.match(text)
		if match:
			num = match.group(1)
			if num in questions:
				issues.append(f"duplicate question {num}")
				active = questions[num]
				continue
			stem, options = _split_flk2_inline_options(match.group(2).strip())
			active = {"num": num, "stem_parts": [], "options": options}
			if stem:
				active["stem_parts"].append(stem)
			questions[num] = active
			question_order.append(num)
			continue
		if active is None:
			continue
		stem, options = _split_flk2_inline_options(text, allow_non_a_start=bool(active["options"]))
		if options:
			if stem and not active["options"]:
				active["stem_parts"].append(stem)
			active["options"].extend(options)
		elif not active["options"]:
			active["stem_parts"].append(text)
		elif active["options"]:
			active["options"][-1]["text"] = (
				active["options"][-1]["text"] + " " + text
			).strip()

	for num in question_order:
		questions[num]["question"] = " ".join(questions[num].pop("stem_parts", [])).strip()
	return questions, question_order, issues


def _parse_flk2_answers(lines: list[str], question_order: list[str]) -> tuple[dict, list[str]]:
	answers: dict[str, dict] = {}
	issues: list[str] = []
	current_num: str | None = None
	bare_cursor = 0
	reason_mode = False
	reason_cursor = 0

	def next_unassigned() -> str | None:
		return next((num for num in question_order if num not in answers), None)

	for text in lines:
		if not text:
			continue
		if FLK2_FINAL_ANSWER_RE.match(text):
			continue
		if re.match(
			r"^(?:self-assessment|scoring guide|mcq test score|chapter\s+summary|"
			r"chapter\s+.*(?:self-assessment|scoring))",
			text,
			re.I,
		):
			break
		if re.match(r"^reasons?:?$", text, re.I):
			reason_mode = True
			current_num = None
			continue

		match = FLK2_EXPLICIT_ANSWER_RE.match(text)
		if match:
			num, letter, tail = match.group(1), match.group(2), match.group(3)
		elif (match := FLK2_ANSWER_RE.match(text)):
			num, letter, tail = match.group(1), match.group(2), match.group(3)
		elif (match := FLK2_NUMBERED_ANSWER_RE.match(text)):
			num, letter, tail = match.group(1), match.group(2), match.group(3)
		elif (match := FLK2_BARE_ANSWER_RE.match(text)):
			num = next_unassigned()
			letter, tail = match.group(1), match.group(2)
			bare_cursor += 1
		elif (match := FLK2_BARE_LETTER_RE.match(text)):
			num = next_unassigned()
			letter, tail = match.group(1), ""
			bare_cursor += 1
		else:
			correct = re.match(r"^Correct\s+Answer\s*:?\s*([A-E])(?:\s*[-–:]\s*|\s+)?(.*)$", text, re.I)
			if not correct:
				if reason_mode and reason_cursor < len(question_order):
					num = question_order[reason_cursor]
					_append_verified_explanation(answers[num], text)
					reason_cursor += 1
					current_num = num
					continue
				if current_num:
					_append_verified_explanation(answers[current_num], text)
				continue
				num = next_unassigned()
				letter, tail = correct.group(1), correct.group(2)

		if not num:
			issues.append(f"answer {letter} has no question to map to")
			continue
		if num in answers:
			issues.append(f"duplicate answer {num}")
			continue
		answers[num] = {"correct_letter": letter.upper(), "explanation": ""}
		current_num = num
		if tail.strip():
			_append_verified_explanation(answers[num], _strip_answer_lead(tail))
			if re.match(r"^\s*[\(\[]?\s*or\s+(?:possibly\s+)?[A-E]\b", tail, re.I):
				issues.append(f"ambiguous answer marker for question {num}: {text}")

	return answers, issues


def _strip_answer_lead(text: str) -> str:
	return re.sub(r"^(?:Explanation|Reason)\s*:\s*", "", text or "", flags=re.I).strip()


def _append_verified_explanation(row: dict, text: str) -> None:
	text = _strip_answer_lead(text)
	if text:
		row["explanation"] = " ".join(filter(None, (row.get("explanation"), text))).strip()


def _validate_flk2_section(questions: dict, answers: dict, question_order: list[str]) -> list[str]:
	issues: list[str] = []
	question_nums = set(question_order)
	answer_nums = set(answers)
	for num in sorted(question_nums - answer_nums, key=_numeric_key):
		issues.append(f"missing answer for question {num}")
	for num in sorted(answer_nums - question_nums, key=_numeric_key):
		issues.append(f"answer has no matching question {num}")
	for num in question_order:
		question = questions[num]
		letters = [option.get("letter") for option in question.get("options", [])]
		if letters != ["A", "B", "C", "D", "E"]:
			issues.append(f"question {num} options are not exactly A-E: {letters}")
		if not question.get("question"):
			issues.append(f"question {num} has an empty stem")
		answer = answers.get(num)
		if not answer:
			continue
		if answer["correct_letter"] not in letters:
			issues.append(f"question {num} answer {answer['correct_letter']} is not an option")
		if not answer.get("explanation", "").strip():
			issues.append(f"question {num} has no explanation")
	return issues


def _numeric_key(value: str) -> tuple[int, ...]:
	return tuple(int(part) for part in value.split("."))


def audit_flk2_mcqs() -> dict:
	"""Read-only FLK2 MCQ reconciliation report without question-body output."""
	from aimaticlearning.lms_learning.course_upload import audit_flk2_sources

	source_audit = audit_flk2_sources()
	results = []
	for source in source_audit.get("subjects", []):
		path = _resolve_source_path(source["source_file"])
		parsed = parse_verified_flk2_mcqs_from_docx(path)
		results.append(
			{
				"subject": parsed["subject"],
				"source_file": parsed["source_file"],
				"ready": parsed["ready"],
				"question_count": parsed["question_count"],
				"sections": parsed["sections"],
				"issues": parsed["issues"],
			}
		)
	return {
		"ready": source_audit["ready"] and all(row["ready"] for row in results),
		"source_issues": source_audit["issues"],
		"subjects": results,
		"publishable_subjects": [row["subject"] for row in results if row["ready"]],
		"blocked_subjects": [row["subject"] for row in results if not row["ready"]],
		"total_reconciled_questions": sum(row["question_count"] for row in results),
	}


def publish_flk2_verified_mcqs(subjects: list[str] | None = None) -> dict:
	"""Publish only exact source-mapped FLK2 chapter quizzes; leave blocked subjects untouched."""
	from aimaticlearning.lms_learning.course_upload import audit_flk2_sources
	from aimaticlearning.lms_learning.outline_sync import ensure_quiz_lesson

	source_audit = audit_flk2_sources()
	if not source_audit["ready"]:
		frappe.throw("FLK2 source preflight failed: " + "; ".join(source_audit["issues"]))

	parsed_by_subject = {}
	for source in source_audit["subjects"]:
		parsed = parse_verified_flk2_mcqs_from_docx(_resolve_source_path(source["source_file"]))
		parsed_by_subject[parsed["subject"]] = parsed

	requested = set(subjects or parsed_by_subject)
	unknown = sorted(requested - set(parsed_by_subject))
	if unknown:
		frappe.throw("Unknown FLK2 subjects: " + ", ".join(unknown))
	blocked = sorted(subject for subject in requested if not parsed_by_subject[subject]["ready"])
	if blocked:
		frappe.throw(
			"Refusing to publish subjects with unresolved answer mappings: "
			+ "; ".join(blocked)
		)

	results = []
	for subject in sorted(requested):
		parsed = parsed_by_subject[subject]
		course_name = frappe.db.get_value("LMS Course", {"title": subject}, "name")
		if not course_name:
			frappe.throw(f"Missing FLK2 course: {subject}")
		course = frappe.get_doc("LMS Course", course_name)
		if not course.published:
			frappe.throw(f"FLK2 course is not published: {subject}")
		module_name = frappe.db.get_value(
			"Learning Module Config", {"lms_course": course.name}, "name"
		)
		if not module_name:
			frappe.throw(f"Missing learning module config: {subject}")
		module = frappe.get_doc("Learning Module Config", module_name)
		profiles = frappe.get_all(
			"Learning Chapter Profile",
			filters={"learning_module": module.name},
			fields=["name", "course_chapter", "chapter_title", "chapter_quiz"],
			order_by="creation asc",
		)
		if len(profiles) != len(parsed["sections"]):
			frappe.throw(
				f"Chapter profile count mismatch for {subject}: "
				f"{len(profiles)} profiles, {len(parsed['sections'])} source sections"
			)

		questions_by_profile: dict[str, list[str]] = defaultdict(list)
		items_by_section: dict[int, list[dict]] = defaultdict(list)
		for item in parsed["questions"]:
			items_by_section[item["section_id"]].append(item)
		for section_id, profile_row in enumerate(profiles, start=1):
			items = items_by_section[section_id]
			if len(items) != parsed["sections"][section_id - 1]["question_count"]:
				frappe.throw(f"Question count mismatch for {subject} section {section_id}")
			for item in items:
				question_name = _upsert_lms_question_by_source(
					{
						"question": item["question"],
						"options": item["options"],
						"concept": profile_row.chapter_title,
						"source_reference": item["source_reference"],
					}
				)
				if not question_name:
					frappe.throw(f"Question upsert failed: {item['source_reference']}")
				meta_name = frappe.db.get_value(
					"Learning Question Meta", {"source_reference": item["source_reference"]}, "name"
				)
				meta_fields = {
					"doctype": "Learning Question Meta",
					"lms_question": question_name,
					"learning_module": module.name,
					"course_chapter": profile_row.course_chapter,
					"concept": profile_row.chapter_title,
					"difficulty": "Medium",
					"question_role": "Chapter MCQ",
					"source_reference": item["source_reference"],
					"revision": 1,
					"ai_generated": 0,
				}
				if meta_name:
					meta = frappe.get_doc("Learning Question Meta", meta_name)
					meta.update(meta_fields)
					meta.save(ignore_permissions=True)
				else:
					frappe.get_doc(meta_fields).insert(ignore_permissions=True)
				questions_by_profile[profile_row.name].append(question_name)

		published_quizzes = []
		for profile_row in profiles:
			profile = frappe.get_doc("Learning Chapter Profile", profile_row.name)
			quiz_name = profile.chapter_quiz or _ensure_chapter_quiz(
				course, profile.course_chapter, profile.chapter_title
			)
			profile.chapter_quiz = quiz_name
			profile.save(ignore_permissions=True)
			quiz = frappe.get_doc("LMS Quiz", quiz_name)
			quiz.set("questions", [])
			for question_name in questions_by_profile[profile.name]:
				quiz.append("questions", {"question": question_name, "marks": 1})
			quiz.max_attempts = 0
			quiz.show_answers = 1
			quiz.passing_percentage = 60
			quiz.total_marks = len(questions_by_profile[profile.name])
			quiz.save(ignore_permissions=True)
			ensure_quiz_lesson(profile.name)
			profile.refresh_mcq_count()
			published_quizzes.append(quiz.name)

		question_count = sum(len(names) for names in questions_by_profile.values())
		module.module_assessment_count = 0
		module.import_status = "Published"
		module.save(ignore_permissions=True)
		results.append(
			{
				"subject": subject,
				"course": course.name,
				"learning_module": module.name,
				"questions": question_count,
				"chapters": len(profiles),
				"quizzes": published_quizzes,
				"module_assessment_created": False,
			}
		)

	frappe.db.commit()
	return {
		"mode": "flk2_verified_chapter_mcqs",
		"published_to_students": True,
		"results": results,
		"blocked_subjects": [
			subject for subject, parsed in parsed_by_subject.items() if not parsed["ready"]
		],
		"total_published_questions": sum(row["questions"] for row in results),
		"module_assessment_created": False,
		"flashcards_published": False,
		"enrolments_created": False,
	}


def _without_explanation_label(text: str) -> str:
	match = re.search(r"(?:Explanation|Reason)\s*:\s*", text or "", re.I)
	return (text[match.end() :] if match else text or "").strip()


def _without_learner_source_trace(text: str) -> str:
	"""Keep internal source metadata out of student-facing MCQ explanations."""
	return LEARNER_SOURCE_TRAIL_RE.sub("", text or "").strip()


def _repeated_choice_explanation_updates(row: dict) -> dict:
	"""Keep a shared general explanation only under the correct answer choice."""
	explained = [
		index
		for index in range(1, 6)
		if (row.get(f"option_{index}") or "").strip()
		and (row.get(f"explanation_{index}") or "").strip()
	]
	if len(explained) < 2:
		return {}
	explanations = {(row.get(f"explanation_{index}") or "").strip() for index in explained}
	if len(explanations) != 1:
		return {}
	correct = [index for index in explained if cint(row.get(f"is_correct_{index}"))]
	if len(correct) != 1:
		return {}
	return {
		f"explanation_{index}": ""
		for index in explained
		if index != correct[0]
	}


def _append_explanation(row: dict, text: str) -> None:
	text = (text or "").strip()
	if not text:
		return
	parts = [row.get("explanation") or row.get("answer_summary") or "", text]
	row["explanation"] = " ".join(part for part in parts if part).strip()


def _split_options(text: str) -> list[dict]:
	options: list[dict] = []
	for part in OPTION_SPLIT_RE.split(text.strip()):
		part = part.strip()
		if not part:
			continue
		match = OPTION_RE.match(part)
		if match:
			options.append({"letter": match.group(1).upper(), "text": match.group(2).strip()})
	if not options:
		match = OPTION_RE.match(text.strip())
		if match:
			options.append({"letter": match.group(1).upper(), "text": match.group(2).strip()})
	return options


def _normalize_question_row(row: dict, section_id: int, num: str, source_label: str) -> dict:
	correct_letter = (row.get("correct_letter") or "").upper()
	explanation = row.get("explanation") or row.get("answer_summary") or ""
	options = []
	for opt in row["options"][:5]:
		is_correct = opt["letter"] == correct_letter
		options.append(
			{
				"text": opt["text"],
				"is_correct": is_correct,
				"explanation": explanation if is_correct else "",
			}
		)
	return {
		"num": num,
		"section_id": section_id,
		"question": row["question"],
		"options": options,
		"source_reference": f"{source_label} S{section_id} Q{num}",
	}


def _distribute_to_profiles(
	parsed: list[dict], profiles: list, per_chapter_target: int
) -> dict[str, list[dict]]:
	queues: dict[str, list[dict]] = {p.name: [] for p in profiles}
	if not parsed or not profiles:
		return queues

	cap = per_chapter_target
	if len(parsed) < len(profiles) * per_chapter_target:
		cap = max(1, (len(parsed) + len(profiles) - 1) // len(profiles))

	counts: dict[str, int] = defaultdict(int)
	idx = 0
	while idx < len(parsed):
		placed = False
		for profile in profiles:
			if counts[profile.name] < cap:
				queues[profile.name].append(parsed[idx])
				counts[profile.name] += 1
				idx += 1
				placed = True
				break
		if not placed:
			break
	return queues


def import_mcqs_for_module(
	learning_module: str,
	limit: int | None = None,
	source_file: str | None = None,
	per_chapter_target: int | None = None,
	module_assessment_target: int | None = None,
) -> dict:
	module = frappe.get_doc("Learning Module Config", learning_module)
	per_chapter_target = int(
		per_chapter_target or module.target_chapter_mcq_count or DEFAULT_PER_CHAPTER_TARGET
	)
	module_assessment_target = int(
		module_assessment_target or module.module_assessment_count or DEFAULT_MODULE_ASSESSMENT_TARGET
	)

	path = _resolve_source_path(source_file or module.source_file)
	parsed = parse_mcqs_from_docx(path, limit=limit)
	if not parsed:
		frappe.throw(_("No MCQs found in source document."))

	profiles = frappe.get_all(
		"Learning Chapter Profile",
		filters={"learning_module": learning_module},
		fields=["name", "course_chapter", "chapter_title", "chapter_quiz"],
		order_by="creation asc",
	)
	if not profiles:
		frappe.throw(_("No chapter profiles found. Run structure import first."))

	for profile in profiles:
		if profile.chapter_quiz:
			quiz = frappe.get_doc("LMS Quiz", profile.chapter_quiz)
			quiz.set("questions", [])
			quiz.total_marks = 0
			quiz.save(ignore_permissions=True)

	profile_queues = _distribute_to_profiles(parsed, profiles, per_chapter_target)
	created = 0
	linked_chapter = 0
	all_question_names: list[str] = []

	for profile in profiles:
		for item in profile_queues[profile.name]:
			payload = {
				"question": item["question"],
				"options": item["options"],
				"concept": profile.chapter_title,
				"source_reference": item["source_reference"],
			}
			try:
				question_name = _upsert_lms_question_by_source(payload)
			except Exception as exc:
				frappe.log_error(
					title="BLP MCQ import skipped question",
					message=f"{item['source_reference']}: {exc}",
				)
				continue
			if not question_name:
				continue
			all_question_names.append(question_name)

			meta_name = frappe.db.get_value("Learning Question Meta", {"lms_question": question_name})
			meta_fields = {
				"doctype": "Learning Question Meta",
				"lms_question": question_name,
				"learning_module": learning_module,
				"course_chapter": profile.course_chapter,
				"concept": profile.chapter_title,
				"difficulty": "Medium",
				"question_role": "Chapter MCQ",
				"source_reference": item["source_reference"],
				"revision": 1,
			}
			if meta_name:
				frappe.get_doc("Learning Question Meta", meta_name).update(meta_fields).save(
					ignore_permissions=True
				)
			else:
				frappe.get_doc(meta_fields).insert(ignore_permissions=True)
				created += 1

			if profile.chapter_quiz:
				_link_question_to_quiz(profile.chapter_quiz, question_name)
				linked_chapter += 1

		if profile.chapter_quiz:
			quiz = frappe.get_doc("LMS Quiz", profile.chapter_quiz)
			quiz.total_marks = len(quiz.questions or [])
			quiz.save(ignore_permissions=True)

	# Module assessment: up to 150 from full imported pool (even spread).
	assessment_pool = list(dict.fromkeys(all_question_names))
	assessment_count = min(module_assessment_target, len(assessment_pool))
	assessment_selected = _even_select(assessment_pool, assessment_count)

	assessment_quiz = module.module_assessment_quiz
	if assessment_quiz:
		quiz = frappe.get_doc("LMS Quiz", assessment_quiz)
		quiz.set("questions", [])
		for question_name in assessment_selected:
			quiz.append("questions", {"question": question_name, "marks": 1})
			meta_name = frappe.db.get_value("Learning Question Meta", {"lms_question": question_name})
			if meta_name:
				meta = frappe.get_doc("Learning Question Meta", meta_name)
				if meta.question_role == "Chapter MCQ":
					meta.question_role = "Both"
				meta.save(ignore_permissions=True)
		quiz.total_marks = len(assessment_selected)
		quiz.save(ignore_permissions=True)

	module.target_chapter_mcq_count = per_chapter_target
	module.module_assessment_count = module_assessment_target
	module.module_mcq_count = len(assessment_selected)
	module.import_status = "Content Ready"
	module.save(ignore_permissions=True)

	for profile in profiles:
		frappe.get_doc("Learning Chapter Profile", profile.name).refresh_mcq_count()

	frappe.db.commit()
	return {
		"learning_module": learning_module,
		"parsed_from_source": len(parsed),
		"imported_to_chapters": len(all_question_names),
		"per_chapter_target": per_chapter_target,
		"module_assessment_selected": len(assessment_selected),
		"module_assessment_target": module_assessment_target,
		"created_meta": created,
		"linked_chapter_quizzes": linked_chapter,
		"chapter_distribution": {
			p.name: len(profile_queues[p.name]) for p in profiles if profile_queues[p.name]
		},
	}


def import_review_mcqs_for_module(learning_module: str, source_file: str) -> dict:
	module = frappe.get_doc("Learning Module Config", learning_module)
	path = _resolve_source_path(source_file)
	parsed = parse_mcqs_from_docx(path)
	if not parsed:
		frappe.throw(_("No MCQs found in source document."))

	profiles = frappe.get_all(
		"Learning Chapter Profile",
		filters={"learning_module": learning_module},
		fields=["name", "course_chapter", "chapter_title"],
		order_by="creation asc",
	)
	if not profiles:
		frappe.throw(_("No chapter profiles found. Run structure import first."))

	per_chapter_target = max(1, (len(parsed) + len(profiles) - 1) // len(profiles))
	profile_queues = _distribute_to_profiles(parsed, profiles, per_chapter_target)
	created = 0
	all_question_names: list[str] = []

	for profile in profiles:
		for item in profile_queues[profile.name]:
			payload = {
				"question": item["question"],
				"options": item["options"],
				"concept": profile.chapter_title,
				"source_reference": item["source_reference"],
			}
			question_name = _upsert_lms_question_by_source(payload)
			if not question_name:
				continue
			all_question_names.append(question_name)

			meta_name = frappe.db.get_value("Learning Question Meta", {"lms_question": question_name})
			meta_fields = {
				"doctype": "Learning Question Meta",
				"lms_question": question_name,
				"learning_module": learning_module,
				"course_chapter": profile.course_chapter,
				"concept": profile.chapter_title,
				"difficulty": "Medium",
				"question_role": "Chapter MCQ",
				"source_reference": item["source_reference"],
				"revision": 1,
				"ai_generated": 0,
			}
			if meta_name:
				frappe.get_doc("Learning Question Meta", meta_name).update(meta_fields).save(
					ignore_permissions=True
				)
			else:
				frappe.get_doc(meta_fields).insert(ignore_permissions=True)
				created += 1

	module.module_mcq_count = len(all_question_names)
	module.import_status = "Content Ready"
	module.save(ignore_permissions=True)
	frappe.db.commit()
	return {
		"learning_module": learning_module,
		"mode": "review_only",
		"source_file": str(path),
		"parsed_from_source": len(parsed),
		"imported_for_review": len(all_question_names),
		"created_meta": created,
		"chapter_distribution": {
			p.chapter_title: len(profile_queues[p.name]) for p in profiles if profile_queues[p.name]
		},
		"published_to_students": False,
		"linked_to_quizzes": False,
	}


def publish_review_mcqs_for_module(learning_module: str) -> dict:
	module = frappe.get_doc("Learning Module Config", learning_module)
	if not module.lms_course:
		frappe.throw(_("Learning module is not linked to an LMS course."))

	profiles = frappe.get_all(
		"Learning Chapter Profile",
		filters={"learning_module": learning_module},
		fields=["name", "course_chapter", "chapter_title", "chapter_quiz"],
		order_by="creation asc",
	)
	if not profiles:
		frappe.throw(_("No chapter profiles found."))

	questions_by_profile: dict[str, list[dict]] = {}
	for profile in profiles:
		rows = frappe.get_all(
			"Learning Question Meta",
			filters={
				"learning_module": learning_module,
				"course_chapter": profile.course_chapter,
				"question_role": ["in", ["Chapter MCQ", "Both"]],
			},
			fields=["lms_question", "source_reference"],
			order_by="creation asc",
		)
		if not rows:
			frappe.throw(
				_("No reviewed MCQs found for chapter {0}.").format(profile.chapter_title)
			)
		questions_by_profile[profile.name] = rows

	course = frappe.get_doc("LMS Course", module.lms_course)
	linked_questions = 0
	quiz_lessons = 0
	quiz_names: list[str] = []

	for profile_row in profiles:
		profile = frappe.get_doc("Learning Chapter Profile", profile_row.name)
		quiz_name = profile.chapter_quiz or _ensure_chapter_quiz(
			course, profile.course_chapter, profile.chapter_title
		)
		if profile.chapter_quiz != quiz_name:
			profile.chapter_quiz = quiz_name
			profile.save(ignore_permissions=True)

		quiz = frappe.get_doc("LMS Quiz", quiz_name)
		quiz.set("questions", [])
		for row in questions_by_profile[profile.name]:
			quiz.append("questions", {"question": row.lms_question, "marks": 1})
		quiz.max_attempts = 0
		quiz.show_answers = 1
		quiz.passing_percentage = 60
		quiz.total_marks = len(questions_by_profile[profile.name])
		quiz.save(ignore_permissions=True)
		linked_questions += len(questions_by_profile[profile.name])
		quiz_names.append(quiz.name)

		if sync_profile_outline(profile.name).get("quiz_lesson"):
			quiz_lessons += 1

	module.import_status = "Published"
	module.save(ignore_permissions=True)
	frappe.db.commit()
	return {
		"learning_module": learning_module,
		"course": module.lms_course,
		"published_to_students": True,
		"linked_to_quizzes": True,
		"chapters_published": len(profiles),
		"quizzes": quiz_names,
		"linked_questions": linked_questions,
		"quiz_lessons": quiz_lessons,
		"module_assessment_created": False,
	}


def _even_select(pool: list[str], count: int) -> list[str]:
	if count >= len(pool):
		return list(pool)
	if count <= 0:
		return []
	step = len(pool) / count
	selected: list[str] = []
	for i in range(count):
		idx = min(int(i * step), len(pool) - 1)
		if pool[idx] not in selected:
			selected.append(pool[idx])
	for name in pool:
		if len(selected) >= count:
			break
		if name not in selected:
			selected.append(name)
	return selected[:count]


def _resolve_source_path(source_file: str | None) -> Path:
	if not source_file:
		frappe.throw(_("Source file is required."))

	if frappe.db.exists("File", source_file):
		file_url = frappe.db.get_value("File", source_file, "file_url")
		if file_url:
			source_file = file_url

	if source_file.startswith("/private/files/"):
		relative = source_file.split("/private/files/", 1)[1]
		return Path(get_site_path("private", "files", relative))
	if source_file.startswith("/files/"):
		relative = source_file.split("/files/", 1)[1]
		return Path(get_site_path("public", "files", relative))
	if Path(source_file).exists():
		return Path(source_file)
	raise frappe.ValidationError(_("Could not resolve source file path."))


def _upsert_lms_question_by_source(item: dict) -> str | None:
	if not any(opt.get("is_correct") for opt in item.get("options") or []):
		return None

	source_reference = item.get("source_reference")
	existing = None
	if source_reference:
		existing = frappe.db.get_value(
			"Learning Question Meta", {"source_reference": source_reference}, "lms_question"
		)

	fields = {
		"doctype": "LMS Question",
		"question": item.get("question"),
		"type": "Choices",
		"multiple": 0,
	}
	for idx in range(1, 6):
		fields[f"option_{idx}"] = ""
		fields[f"is_correct_{idx}"] = 0
		fields[f"explanation_{idx}"] = ""

	for idx, opt in enumerate(item.get("options") or [], start=1):
		if idx > 5:
			break
		fields[f"option_{idx}"] = opt.get("text")
		fields[f"is_correct_{idx}"] = 1 if opt.get("is_correct") else 0
		if opt.get("explanation"):
			fields[f"explanation_{idx}"] = _without_learner_source_trace(
				opt.get("explanation")
			)

	if existing:
		doc = frappe.get_doc("LMS Question", existing)
		doc.update(fields)
		doc.save(ignore_permissions=True)
		return doc.name

	doc = frappe.get_doc(fields)
	doc.insert(ignore_permissions=True)
	return doc.name


def remove_learner_source_traces() -> dict:
	"""Remove legacy filesystem source suffixes from existing LMS question explanations."""
	fields = [f"explanation_{index}" for index in range(1, 6)]
	rows = frappe.db.get_all("LMS Question", fields=["name", *fields])
	updated_questions = 0
	updated_explanations = 0
	for row in rows:
		updates = {}
		for field in fields:
			original = row.get(field) or ""
			cleaned = _without_learner_source_trace(original)
			if cleaned != original:
				updates[field] = cleaned
		if not updates:
			continue
		frappe.db.set_value("LMS Question", row.name, updates, update_modified=False)
		updated_questions += 1
		updated_explanations += len(updates)
	frappe.db.commit()
	return {
		"updated_questions": updated_questions,
		"updated_explanations": updated_explanations,
	}


def remove_repeated_choice_explanations() -> dict:
	"""Remove legacy duplicate feedback while retaining it beneath the correct option."""
	fields = [
		"name",
		*[f"option_{index}" for index in range(1, 6)],
		*[f"is_correct_{index}" for index in range(1, 6)],
		*[f"explanation_{index}" for index in range(1, 6)],
	]
	rows = frappe.db.get_all("LMS Question", fields=fields)
	updated_questions = 0
	updated_explanations = 0
	for row in rows:
		updates = _repeated_choice_explanation_updates(row)
		if not updates:
			continue
		frappe.db.set_value("LMS Question", row.name, updates, update_modified=False)
		updated_questions += 1
		updated_explanations += len(updates)
	frappe.db.commit()
	return {
		"updated_questions": updated_questions,
		"updated_explanations": updated_explanations,
	}
