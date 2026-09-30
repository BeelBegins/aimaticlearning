import tempfile
import unittest
from pathlib import Path

from docx import Document

from aimaticlearning.lms_learning.content_generation import _normalise_source_text
from aimaticlearning.lms_learning.mcq_import import (
	_parse_flk2_answers,
	_repeated_choice_explanation_updates,
	_split_flk2_inline_options,
	_without_learner_source_trace,
	parse_mcqs_from_docx,
)


class TestContentGrounding(unittest.TestCase):
	def test_normalise_source_text_removes_html_and_collapses_whitespace(self):
		value = "<p>UK&nbsp;rule</p><p>Second line</p>"
		self.assertEqual(_normalise_source_text(value, strip_html=True), "UK rule Second line")

	def test_student_explanation_omits_filesystem_source_trace(self):
		value = (
			"The actus reus requires voluntary conduct. "
			"Source: /home/nabeel/.codex/attachments/example/notes.docx, Chapter 1."
		)
		self.assertEqual(
			_without_learner_source_trace(value),
			"The actus reus requires voluntary conduct.",
		)

	def test_repeated_choice_explanation_is_retained_only_for_correct_option(self):
		question = {
			**{f"option_{index}": f"Option {index}" for index in range(1, 6)},
			**{f"explanation_{index}": "One general explanation." for index in range(1, 6)},
			**{f"is_correct_{index}": 1 if index == 2 else 0 for index in range(1, 6)},
		}
		self.assertEqual(
			_repeated_choice_explanation_updates(question),
			{
				"explanation_1": "",
				"explanation_3": "",
				"explanation_4": "",
				"explanation_5": "",
			},
		)

	def test_mcq_parser_keeps_detailed_explanation(self):
		document = Document()
		for paragraph in (
			"MCQ Practice",
			"Question 1 Marcus owes £18,500 to various creditors.",
			"A. Petition for bankruptcy",
			"B. Apply for a Debt Relief Order",
			"C. Propose an Individual Voluntary Arrangement",
			"D. Wait until the debts exceed £20,000",
			"ANSWER KEY WITH DETAILED EXPLANATIONS",
			"Question 1: Answer B - Apply for a Debt Relief Order.",
			"Correct Answer: B - Apply for a Debt Relief Order.",
			"Explanation:",
			"A Debt Relief Order is designed for debtors with nominal assets and minimal income.",
			"It is the simplest and cheapest solution in this scenario.",
		):
			document.add_paragraph(paragraph)

		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "questions.docx"
			document.save(path)
			questions = parse_mcqs_from_docx(path)

		self.assertEqual(len(questions), 1)
		correct = next(option for option in questions[0]["options"] if option["is_correct"])
		self.assertIn("nominal assets", correct["explanation"])
		self.assertIn("simplest and cheapest", correct["explanation"])

	def test_mcq_parser_accepts_decimal_and_inline_answer_keys(self):
		document = Document()
		for paragraph in (
			"Questions 1.1-1.20",
			"Question 1.1 A claimant suffers loss after negligent advice.",
			"A) Contract only B) Tort only C) Both contract and tort D) Neither",
			"ANSWER KEY - CHAPTER 1",
			"Answer 1.1: C Reason: Concurrent liability may arise on the same facts.",
			"Multiple Choice Questions",
			"Question 1 A solicitor must act fairly.",
			"A) Ignore the duty B) Follow the rules C) Delay the case D) Withdraw",
			"Answer Key",
			"Question 1: B Reason: The professional duty requires compliance.",
		):
			document.add_paragraph(paragraph)

		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "mixed-questions.docx"
			document.save(path)
			questions = parse_mcqs_from_docx(path)

		self.assertEqual(len(questions), 2)
		self.assertIn("Concurrent liability", questions[0]["options"][2]["explanation"])
		self.assertIn("professional duty", questions[1]["options"][1]["explanation"])

	def test_mcq_parser_accepts_correct_answer_with_en_dash_and_reason(self):
		document = Document()
		for paragraph in (
			"LEGAL SERVICES",
			"SQE1-Style Multiple Choice Questions with Answers and Explanations",
			"Questions 1–20",
			"Question 1. Which professional is regulated by the Bar Council?",
			"A. A solicitor",
			"B. A barrister",
			"C. A notary",
			"D. A paralegal",
			"E. A legal executive",
			"Answers and Explanations",
			"Question 1: Correct answer – B",
			"Reason: Barristers are regulated by the Bar Council.",
		):
			document.add_paragraph(paragraph)

		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "MCQs.docx"
			document.save(path)
			questions = parse_mcqs_from_docx(path)

		self.assertEqual(len(questions), 1)
		self.assertEqual(questions[0]["section_id"], 1)
		self.assertFalse(questions[0]["question"].startswith("."))
		correct = next(option for option in questions[0]["options"] if option["is_correct"])
		self.assertEqual(correct["text"], "A barrister")
		self.assertIn("regulated by the Bar Council", correct["explanation"])

	def test_flk2_helpers_reconcile_bare_answers_and_embedded_option_text(self):
		prefix, options = _split_flk2_inline_options(
			"Client B. must be protected. A) Alpha B) Beta C) Gamma D) Delta E) Epsilon"
		)
		self.assertEqual(prefix, "Client B. must be protected.")
		self.assertEqual([option["letter"] for option in options], ["A", "B", "C", "D", "E"])

		answers, issues = _parse_flk2_answers(
			["B", "C", "REASONS:", "First source reason.", "Second source reason."],
			["1", "2"],
		)
		self.assertEqual(issues, [])
		self.assertEqual(answers["1"]["correct_letter"], "B")
		self.assertIn("First source reason", answers["1"]["explanation"])
		self.assertEqual(answers["2"]["correct_letter"], "C")
		self.assertIn("Second source reason", answers["2"]["explanation"])


if __name__ == "__main__":
	unittest.main()
